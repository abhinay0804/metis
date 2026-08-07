import json
import csv
from docx import Document
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
import os

os.makedirs('test_files', exist_ok=True)

# 1. TXT
with open('test_files/sample.txt', 'w') as f:
    f.write("Patient Record\nName: John Doe\nEmail: john.doe@acme.com\nSSN: 123-45-6789\nPhone: (555) 123-4567\n")

# 2. JSON
with open('test_files/sample.json', 'w') as f:
    json.dump({
        "patient_record": {
            "name": "Jane Smith",
            "email": "jane.smith@hospital.org",
            "ssn": "987-65-4321",
            "phone": "555-987-6543"
        }
    }, f)

# 3. CSV
with open('test_files/sample.csv', 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['Name', 'Email', 'SSN', 'Phone'])
    writer.writerow(['Alice Williams', 'alice@corp.com', '111-22-3333', '123-456-7890'])

# 4. DOCX
doc = Document()
doc.add_heading('Employee Data', 0)
doc.add_paragraph('Employee Name: Bob Brown')
doc.add_paragraph('Contact: bob.brown@enterprise.net')
doc.add_paragraph('SSN: 444-55-6666')
doc.add_paragraph('Phone: (999) 888-7777')
doc.save('test_files/sample.docx')

# 5. PDF
c = canvas.Canvas('test_files/sample.pdf')
c.drawString(100, 750, "Confidential Client Info")
c.drawString(100, 730, "Name: Carol Davis")
c.drawString(100, 710, "Email: carol.d@startup.io")
c.drawString(100, 690, "SSN: 777-88-9999")
c.drawString(100, 670, "Phone: 222-333-4444")
c.save()

# 6. XLSX
wb = Workbook()
ws = wb.active
ws.append(['Name', 'Email', 'SSN', 'Phone'])
ws.append(['David Miller', 'david.m@bank.com', '000-11-2222', '444-555-6666'])
wb.save('test_files/sample.xlsx')

# 7. PNG
img = Image.new('RGB', (400, 200), color=(255, 255, 255))
d = ImageDraw.Draw(img)
text = "Scanned ID Card\nName: Eve Wilson\nEmail: eve.w@agency.gov\nSSN: 555-66-7777\nPhone: (111) 222-3333"
d.text((10,10), text, fill=(0,0,0))
img.save('test_files/sample.png')

print("Generated all test files in 'test_files' directory.")
